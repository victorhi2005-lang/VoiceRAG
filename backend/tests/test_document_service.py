import importlib.util
import io
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock


BACKEND_DIR = Path(__file__).resolve().parents[1]


def load_document_service():
    db_stub = types.ModuleType("db")
    for name in (
        "get_all_source_filenames",
        "get_source_file_info",
        "get_source_filenames",
        "insert_source_record",
    ):
        setattr(db_stub, name, mock.Mock())
    db_stub.get_source_filenames.return_value = []
    db_stub.get_all_source_filenames.return_value = []

    rag_stub = types.ModuleType("rag_service")
    rag_stub.delete_source_from_index = mock.Mock()
    rag_stub.index_source_document = mock.Mock(return_value=1)
    rag_stub.should_use_deep_analysis = mock.Mock(return_value=False)

    with mock.patch.dict(sys.modules, {"db": db_stub, "rag_service": rag_stub}):
        spec = importlib.util.spec_from_file_location(
            "document_service_under_test",
            BACKEND_DIR / "document_service.py",
        )
        module = importlib.util.module_from_spec(spec)
        assert spec and spec.loader
        spec.loader.exec_module(module)
        return module


document_service = load_document_service()


class FakeUpload:
    def __init__(self, data: bytes, filename: str, content_type: str):
        self.file = io.BytesIO(data)
        self.filename = filename
        self.content_type = content_type


class DocumentServiceTests(unittest.TestCase):
    def test_pdf_is_delivered_inline(self):
        media_type, disposition = document_service.get_document_delivery_options(
            "sample.pdf",
            "application/pdf",
        )
        self.assertEqual(media_type, "application/pdf")
        self.assertEqual(disposition, "inline")

    def test_docx_is_delivered_as_attachment(self):
        media_type, disposition = document_service.get_document_delivery_options(
            "sample.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertEqual(
            media_type,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertEqual(disposition, "attachment")

    def test_txt_supports_big5_and_line_locations(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "筆記.txt"
            path.write_bytes("第一行\n第二行".encode("big5"))

            segments = document_service.extract_text_segments(path)
            media_type, disposition = document_service.get_document_delivery_options(
                path.name,
                "text/plain",
                path,
            )

        self.assertEqual(segments[0]["location_type"], "lines")
        self.assertEqual(segments[0]["line_start"], 1)
        self.assertIn("第一行", segments[0]["text"])
        self.assertEqual(media_type, "text/plain; charset=big5")
        self.assertEqual(disposition, "inline")

    def test_markdown_headings_become_location_labels(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notes.md"
            path.write_text("# 第一章\n內容 A\n## 第二章\n內容 B", encoding="utf-8")

            segments = document_service.extract_text_segments(path)

        self.assertEqual([segment["location_label"] for segment in segments], ["第一章", "第二章"])
        self.assertEqual(segments[1]["section_title"], "第二章")

    def test_pdf_without_signature_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "fake.pdf"
            path.write_bytes(b"not a pdf")
            with self.assertRaisesRegex(document_service.DocumentProcessingError, "有效的 PDF"):
                document_service.validate_document_upload(path, "application/pdf")

    def test_pdf_pages_and_tables_keep_page_number(self):
        class FakePage:
            def extract_text(self):
                return "頁面文字"

            def extract_tables(self):
                return [[["欄一", "欄二"], ["值一", "值二"]]]

        class FakePdf:
            pages = [FakePage()]

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

        pdfplumber_stub = types.ModuleType("pdfplumber")
        pdfplumber_stub.open = mock.Mock(return_value=FakePdf())
        with mock.patch.dict(sys.modules, {"pdfplumber": pdfplumber_stub}):
            segments = document_service.extract_pdf_segments(Path("sample.pdf"))

        self.assertEqual(segments[0]["page_number"], 1)
        self.assertIn("欄一 | 欄二", segments[0]["text"])

    def test_pdf_without_extractable_text_reports_ocr_limitation(self):
        class EmptyPage:
            def extract_text(self):
                return ""

            def extract_tables(self):
                return []

        class EmptyPdf:
            pages = [EmptyPage()]

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

        pdfplumber_stub = types.ModuleType("pdfplumber")
        pdfplumber_stub.open = mock.Mock(return_value=EmptyPdf())
        with (
            mock.patch.dict(sys.modules, {"pdfplumber": pdfplumber_stub}),
            self.assertRaisesRegex(document_service.DocumentProcessingError, "不支援 OCR"),
        ):
            document_service.extract_pdf_segments(Path("scan.pdf"))

    def test_real_pdf_package_extracts_page_text_when_available(self):
        try:
            import pdfplumber  # noqa: F401
        except ImportError:
            self.skipTest("pdfplumber is not installed in this interpreter")

        stream = b"BT /F1 12 Tf 72 720 Td (Hello VoiceRAG) Tj ET"
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
        pdf = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for index, body in enumerate(objects, start=1):
            offsets.append(len(pdf))
            pdf.extend(f"{index} 0 obj\n".encode("ascii"))
            pdf.extend(body)
            pdf.extend(b"\nendobj\n")
        xref_offset = len(pdf)
        pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
        pdf.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        pdf.extend(
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.pdf"
            path.write_bytes(pdf)
            segments = document_service.extract_pdf_segments(path)

        self.assertEqual(segments[0]["page_number"], 1)
        self.assertIn("Hello VoiceRAG", segments[0]["text"])

    def test_docx_keeps_heading_and_table_text(self):
        class FakeStyle:
            def __init__(self, name):
                self.name = name

        class FakeChild:
            def __init__(self, kind, text="", style="", rows=None):
                self.tag = f"}}{kind}"
                self.text = text
                self.style = style
                self.rows = rows or []

        children = [
            FakeChild("p", "第一章", "Heading 1"),
            FakeChild("p", "段落內容", "Normal"),
            FakeChild("tbl", rows=[["欄一", "欄二"], ["值一", "值二"]]),
        ]

        class FakeBody:
            def iterchildren(self):
                return iter(children)

        class FakeDocumentObject:
            element = types.SimpleNamespace(body=FakeBody())

        class FakeParagraph:
            def __init__(self, child, _parent):
                self.text = child.text
                self.style = FakeStyle(child.style)

        class FakeCell:
            def __init__(self, text):
                self.text = text

        class FakeRow:
            def __init__(self, cells):
                self.cells = [FakeCell(cell) for cell in cells]

        class FakeTable:
            def __init__(self, child, _parent):
                self.rows = [FakeRow(row) for row in child.rows]

        docx_module = types.ModuleType("docx")
        docx_module.Document = mock.Mock(return_value=FakeDocumentObject())
        docx_table_module = types.ModuleType("docx.table")
        docx_table_module.Table = FakeTable
        docx_text_module = types.ModuleType("docx.text")
        docx_paragraph_module = types.ModuleType("docx.text.paragraph")
        docx_paragraph_module.Paragraph = FakeParagraph

        with mock.patch.dict(sys.modules, {
            "docx": docx_module,
            "docx.table": docx_table_module,
            "docx.text": docx_text_module,
            "docx.text.paragraph": docx_paragraph_module,
        }):
            segments = document_service.extract_docx_segments(Path("sample.docx"))

        self.assertEqual(segments[0]["location_label"], "第一章")
        self.assertIn("欄一 | 欄二", segments[0]["text"])

    def test_real_docx_package_round_trip_when_available(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed in this interpreter")

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.docx"
            document = Document()
            document.add_heading("專案說明", level=1)
            document.add_paragraph("這是一段繁體中文內容。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "名稱"
            table.cell(0, 1).text = "數值"
            table.cell(1, 0).text = "測試"
            table.cell(1, 1).text = "100"
            document.save(path)

            segments = document_service.extract_docx_segments(path)

        self.assertEqual(segments[0]["location_label"], "專案說明")
        self.assertIn("測試 | 100", segments[0]["text"])

    def test_document_size_limit_removes_partial_file(self):
        class OversizeStream:
            def __init__(self):
                self.remaining = 51

            def read(self, _size):
                if self.remaining <= 0:
                    return b""
                self.remaining -= 1
                return b"x" * (1024 * 1024)

        upload = types.SimpleNamespace(file=OversizeStream())
        with tempfile.TemporaryDirectory() as temp_dir:
            upload_dir = Path(temp_dir)
            with (
                mock.patch.object(document_service, "DOCUMENT_UPLOAD_DIR", upload_dir),
                self.assertRaisesRegex(document_service.DocumentProcessingError, "50 MB"),
            ):
                document_service.save_document_upload(upload, "large.pdf")
            self.assertFalse((upload_dir / "large.pdf").exists())

    def test_failed_document_processing_leaves_no_file_or_index(self):
        upload = FakeUpload(b"# heading", "broken.md", "text/markdown")
        with tempfile.TemporaryDirectory() as temp_dir:
            upload_dir = Path(temp_dir)
            with (
                mock.patch.object(document_service, "DOCUMENT_UPLOAD_DIR", upload_dir),
                mock.patch.object(
                    document_service,
                    "extract_document_segments",
                    side_effect=document_service.DocumentProcessingError("解析失敗"),
                ),
                mock.patch.object(document_service, "delete_source_from_index") as cleanup_mock,
                self.assertRaisesRegex(document_service.DocumentProcessingError, "解析失敗"),
            ):
                document_service.process_document_upload("notebook", upload)

            self.assertEqual(list(upload_dir.iterdir()), [])
            cleanup_mock.assert_called_once()

    def test_markdown_upload_extracts_and_persists_document_metadata(self):
        upload = FakeUpload(
            "# 第一章\n這是文件內容。".encode("utf-8"),
            "notes.md",
            "text/markdown",
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            upload_dir = Path(temp_dir)
            with (
                mock.patch.object(document_service, "DOCUMENT_UPLOAD_DIR", upload_dir),
                mock.patch.object(document_service, "get_all_source_filenames", return_value=[]),
                mock.patch.object(document_service, "index_source_document", return_value=2),
                mock.patch.object(document_service, "insert_source_record", return_value="now") as insert_mock,
            ):
                result = document_service.process_document_upload("notebook", upload)

        self.assertEqual(result["source"]["source_type"], "document")
        self.assertEqual(result["source"]["analysis_status"], "pending")
        self.assertEqual(insert_mock.call_args.kwargs["source_type"], "document")
        self.assertIn("第一章", insert_mock.call_args.kwargs["content_segments_json"])

if __name__ == "__main__":
    unittest.main()
